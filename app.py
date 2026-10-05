import os
import time
import io
import streamlit as st
from google import genai
from docx import Document

st.set_page_config(
    page_title="Banco de Ajustes Razonables - I.PS.I.",
    page_icon="🧠",
    layout="wide",
)

# Cargar automáticamente la clave desde .env si existe
def obtener_clave_maestra():
    if os.path.exists(".env"):
        with open(".env", "r", encoding="utf-8") as f:
            for line in f:
                if line.startswith("GEMINI_API_KEY="):
                    return line.split("=", 1)[1].strip()
    return os.environ.get("GEMINI_API_KEY", "").strip()

ENV_API_KEY = obtener_clave_maestra()

# Escudo del IPSI
if os.path.exists("logo_ipsi.jpg"):
    st.sidebar.image("logo_ipsi.jpg", use_container_width=True)

st.sidebar.markdown("### 🏫 Instituto Psicopedagógico Integral")
st.sidebar.markdown("**Banco Dinámico de Ajustes Razonables (IA)**")
st.sidebar.markdown("---")
st.sidebar.caption("Powered by Hózkar")

# Configuración de Clave API: Si existe clave preconfigurada, usarla por defecto
if ENV_API_KEY:
    st.sidebar.success("✅ **Clave Institucional Activa:** Configuración de cero pasos para los docentes.")
    api_key_input = st.sidebar.text_input(
        "🔑 Clave API (Preconfigurada):", 
        value=ENV_API_KEY,
        type="password", 
        help="Clave configurada para el colegio."
    )
else:
    api_key_input = st.sidebar.text_input(
        "🔑 Clave API de Gemini:", 
        type="password", 
        help="Pega aquí tu clave API personal de Google AI Studio."
    )
    st.sidebar.markdown("[👉 Obtener clave API gratuita en Google AI Studio](https://aistudio.google.com/app/apikey)")

api_key = api_key_input.strip() if api_key_input else ENV_API_KEY

# Selección de modelo
modelo_seleccionado = st.sidebar.selectbox(
    "⚙️ Modelo IA:",
    ["Automático (Recomendado)", "gemini-flash-lite-latest", "gemini-pro-latest", "gemini-2.5-pro"],
    help="El modo Automático reintenta y conmuta de modelo si hay saturación de la API."
)

@st.cache_data
def cargar_opciones_base():
    import pandas as pd
    try:
        file_diag = "listado_de_trastornos_y_deficits_escolares_con_definicion_IPSI_2026.xlsx"
        file_asig = "listado_asignaturas_2026_IPSI.xlsx"
        df_diag_raw = pd.read_excel(file_diag, header=None)
        df_asig_raw = pd.read_excel(file_asig, header=None)
        return df_diag_raw, df_asig_raw
    except Exception as e:
        return None, None

df_diag_raw, df_asig_raw = cargar_opciones_base()

st.title("🧩 Banco Dinámico de Ajustes Razonables y Actividades Sugeridas - I.PS.I.")
st.write("Esta herramienta interactúa en tiempo real con inteligencia artificial para generar planes de apoyo psicopedagógico únicos y contextualizados.")

if df_diag_raw is None or df_asig_raw is None:
    st.error("⚠️ No se encontraron los archivos de Excel requeridos.")
else:
    cursos = ["TR"] + [f"{i}°" for i in range(1, 12)]
    edades = list(range(4, 20))
    generos = ["Masculino", "Femenino"]
    lista_asignaturas = df_asig_raw.iloc[:, 0].dropna().tolist()

    dict_diagnosticos = dict(zip(df_diag_raw.iloc[:, 0].dropna(), df_diag_raw.iloc[:, 1].fillna("Sin descripción disponible.")))
    
    # Agregar la opción de "Sin diagnóstico"
    opciones_diagnostico = ["Ninguno (Sin diagnóstico)"] + list(dict_diagnosticos.keys())
    dict_diagnosticos["Ninguno (Sin diagnóstico)"] = "Estudiante sin un diagnóstico clínico específico."

    with st.form("form_dinamico_ipsi"):
        st.subheader("📋 Parámetros del Caso")
        col1, col2 = st.columns(2)
        with col1:
            curso_sel = st.selectbox("1. CURSO * (Obligatorio)", cursos)
            edad_sel = st.selectbox("2. EDAD * (Obligatorio)", edades)
            genero_sel = st.selectbox("3. GÉNERO * (Obligatorio)", generos)
        with col2:
            asignatura_sel = st.selectbox("4. ASIGNATURA * (Obligatorio)", lista_asignaturas)
            st.markdown("**5. DIAGNÓSTICO(S) * (Obligatorio)**")
            st.caption("💡 *Pase el cursor sobre los nombres para visualizar su definición.*")
            
            diagnosticos_sel = []
            
            with st.container(height=250):
                for diag in opciones_diagnostico:
                    definicion = dict_diagnosticos[diag]
                    if st.checkbox(diag, help=f"Definición: {definicion}"):
                        diagnosticos_sel.append(diag)
            
            st.markdown("---")
            tema_sel = st.text_input("6. TEMA (Opcional - Escriba el tema específico de la clase)")
        
        submitted = st.form_submit_button("✨ Generar Propuesta Inteligente con Gemini")

    if submitted:
        key_limpia = api_key.strip() if api_key else ""
        if not key_limpia:
            st.error("❌ **ERROR:** Se requiere una clave API configurada para generar la propuesta.")
        elif not diagnosticos_sel:
            st.warning("⚠️ Por favor, seleccione al menos un diagnóstico (puede elegir 'Ninguno').")
        else:
            prompt_sistema = f"""
Eres un asesor experto en psicopedagogía e inclusión escolar del Instituto Psicopedagógico Integral (I.PS.I.) de Bogotá.
Tu tarea es generar un plan de apoyo pedagógico detallado, claro y aterrizado al aula de clase para un estudiante con los siguientes parámetros:
- Curso: {curso_sel}
- Edad: {edad_sel} años
- Género: {genero_sel}
- Asignatura: {asignatura_sel}
- Diagnóstico(s): {', '.join(diagnosticos_sel)}
- Tema específico de la clase: {tema_sel if tema_sel else 'No especificado (desarrollar de forma general para la asignatura)'}

Por favor, estructura tu respuesta en las siguientes secciones con lenguaje profesional pero accesible para los docentes:
1. 🛠️ **Sugerencias de Ajustes Razonables** (metodología, apoyos visuales, adaptaciones en tiempos y entorno de aula).
2. 💡 **Actividades Prácticas Sugeridas** (ejercicios concretos y aplicables en la clase de {asignatura_sel}).
3. 📌 **Recomendaciones Clave** (pautas para el docente y articulación con la familia).

Al final de tu respuesta, cierra obligatoriamente con el siguiente crédito exacto:
"BANCO DE AJUSTES RAZONABLES — I.PS.I. | Powered by Hózkar"
"""

            # Definir modelos optimizados
            if modelo_seleccionado == "Automático (Recomendado)":
                modelos_a_probar = ["gemini-flash-lite-latest", "gemini-pro-latest", "gemini-2.5-pro"]
            else:
                modelos_a_probar = [modelo_seleccionado, "gemini-flash-lite-latest", "gemini-pro-latest"]

            exito = False
            texto_resultado = ""
            modelo_usado = ""
            errores_registrados = []

            status_placeholder = st.empty()
            
            with st.spinner("🤖 La IA está analizando el caso y estructurando los ajustes razonables para el I.PS.I..."):
                try:
                    client = genai.Client(api_key=key_limpia)
                except Exception as e_init:
                    st.error(f"❌ Error de inicialización del cliente Gemini: {e_init}")
                    client = None

                if client:
                    for mod in modelos_a_probar:
                        if exito:
                            break
                        
                        max_reintentos = 2
                        for intento in range(1, max_reintentos + 1):
                            try:
                                status_placeholder.info(f"⏳ Consultando a la IA (`{mod}`)...")
                                response = client.models.generate_content(
                                    model=mod,
                                    contents=prompt_sistema,
                                )
                                if response and response.text:
                                    texto_resultado = response.text
                                    modelo_usado = mod
                                    exito = True
                                    break
                            except Exception as e:
                                err_str = str(e)
                                errores_registrados.append(f"[{mod} - Intento {intento}] {err_str}")
                                
                                is_quota_or_server_err = any(k in err_str.upper() for k in ["429", "RESOURCE_EXHAUSTED", "500", "503", "QUOTA", "OVERLOADED"])
                                
                                if is_quota_or_server_err and intento < max_reintentos:
                                    espera = intento * 3
                                    status_placeholder.warning(f"⚠️ Servidor ocupado. Reintentando en {espera}s...")
                                    time.sleep(espera)
                                else:
                                    break

            status_placeholder.empty()

            if exito:
                st.success(f"¡Plan generado exitosamente por la IA!")
                
                # Mostrar el resultado
                st.info("💡 **A continuación se muestra el plan. Puedes copiar el texto directamente o descargarlo usando los botones de abajo.**")
                st.markdown(texto_resultado)
                
                st.markdown("---")
                col_btn1, col_btn2 = st.columns(2)
                
                with col_btn1:
                    st.download_button(
                        label="📥 Descargar como Archivo de Texto (.txt)",
                        data=texto_resultado,
                        file_name=f"Ajustes_{curso_sel}_{asignatura_sel}.txt",
                        mime="text/plain",
                        use_container_width=True
                    )

                with col_btn2:
                    doc = Document()
                    doc.add_heading("Banco de Ajustes Razonables - I.PS.I.", 0)
                    doc.add_paragraph(f"Curso: {curso_sel} | Edad: {edad_sel} | Asignatura: {asignatura_sel}")
                    doc.add_paragraph(f"Diagnósticos: {', '.join(diagnosticos_sel)}")
                    doc.add_paragraph(texto_resultado)
                    
                    bio = io.BytesIO()
                    doc.save(bio)
                    st.download_button(
                        label="📄 Descargar como Documento de Word (.docx)",
                        data=bio.getvalue(),
                        file_name=f"Ajustes_{curso_sel}_{asignatura_sel}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        use_container_width=True
                    )
            else:
                st.error("❌ **No se pudo completar la solicitud con el servicio de IA.**")
                with st.expander("🛠️ Ver detalles del error técnico"):
                    for err in errores_registrados:
                        st.code(err)
